from decimal import Decimal
from io import BytesIO
from unicodedata import combining, normalize

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.modules.members.domain.entities import MemberCobrosRow

COBROS_TITLE = "PLANTILLA DE CARGA DE INFORMACIÓN "
COBROS_SUBTITLE = "CASH COOPERATIVA DE AHORRO Y CRÉDITO ALIANZA DEL VALLE"
COBROS_REFERENCE = "MEMBRESIA COPSSTEC"
IDENTIFICATION_TYPE = "C"
HEADERS = (
    "CONTRAPARTIDA",
    "TIPO DE IDENTIFICACIÓN",
    "IDENTIFICACIÓN",
    "NOMBRE",
    "CORREO",
    "TIPO INSTITUCIÓN",
    "INSTITUCIÓN",
    "TIPO CUENTA",
    "CUENTA",
    "REFERENCIA",
    "MONTO",
    "CÓDIGO",
    "CÓDIGO CUENTA",
)
COLUMN_WIDTHS = (16.71, 16.71, 16.71, 35.86, 35.86, 20.71, 35.86, 20.71, 20.71, 35.86, 20.71, 19.86, 18.57)
PLAN_AMOUNTS: dict[str, Decimal] = {
    "monthly": Decimal("10.00"),
    "quarterly": Decimal("30.00"),
    "semiannual": Decimal("60.00"),
    "annual": Decimal("120.00"),
}
ACCOUNT_TYPE_CODES = {
    "ahorros": ("AHORROS", "02"),
    "ahorro": ("AHORROS", "02"),
    "corriente": ("CORRIENTE", "01"),
}

# label/alias del wizard o nombre CASH → (INSTITUCIÓN, CÓDIGO, TIPO INSTITUCIÓN)
CASH_BANKS: dict[str, tuple[str, int, str]] = {
    "banco pichincha": ("BANCO PICHINCHA", 10, "Banco"),
    "banco guayaquil": ("BANCO DE GUAYAQUIL", 17, "Banco"),
    "banco de guayaquil": ("BANCO DE GUAYAQUIL", 17, "Banco"),
    "banco del pacifico": ("BANCO PACIFICO", 30, "Banco"),
    "banco pacifico": ("BANCO PACIFICO", 30, "Banco"),
    "produbanco": ("BANCO PRODUBANCO", 36, "Banco"),
    "banco produbanco": ("BANCO PRODUBANCO", 36, "Banco"),
    "banco internacional": ("BANCO INTERNACIONAL", 32, "Banco"),
    "banco bolivariano": ("BANCO BOLIVARIANO", 37, "Banco"),
    "banco de machala": ("BANCO DE MACHALA", 25, "Banco"),
    "banco general ruminahui": ("BANCO GENERAL RUMINAHUI", 42, "Banco"),
    "banco del austro": ("BANCO DEL AUSTRO", 35, "Banco"),
    "banco solidario": ("BANCO SOLIDARIO", 59, "Banco"),
    "cooperativa jep": ("COAC JUVENTUD ECUATORIANA PROGRESISTA LTDA.", 213, "Cooperativa"),
    "coop. a. y c. juventud ecuatoriana progresista limitada": (
        "COAC JUVENTUD ECUATORIANA PROGRESISTA LTDA.",
        213,
        "Cooperativa",
    ),
    "coac juventud ecuatoriana progresista ltda.": (
        "COAC JUVENTUD ECUATORIANA PROGRESISTA LTDA.",
        213,
        "Cooperativa",
    ),
    "cooperativa policia nacional": ("COAC POLICIA NACIONAL LTDA.", 9997, "Cooperativa"),
    "coop. a. y c. policia nacional limitada": ("COAC POLICIA NACIONAL LTDA.", 9997, "Cooperativa"),
    "coac policia nacional ltda.": ("COAC POLICIA NACIONAL LTDA.", 9997, "Cooperativa"),
    "banco de machala": ("BANCO DE MACHALA", 25, "Banco"),
}

TITLE_FONT = Font(name="Calibri", size=14, bold=True)
LABEL_FONT = Font(name="Calibri", size=11, bold=True)
DATA_FONT = Font(name="Calibri", size=11)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


def _lookup_key(value: str) -> str:
    decomposed = normalize("NFKD", (value or "").strip().lower())
    folded = "".join(char for char in decomposed if not combining(char))
    return " ".join(folded.split())


def resolve_bank(bank_name: str) -> tuple[str, int | None, str]:
    raw = (bank_name or "").strip()
    if not raw:
        return "", None, ""
    mapped = CASH_BANKS.get(_lookup_key(raw))
    if mapped is None:
        return raw, None, ""
    return mapped


def resolve_account_type(account_type: str) -> tuple[str, str]:
    mapped = ACCOUNT_TYPE_CODES.get(_lookup_key(account_type))
    if mapped is None:
        return (account_type or "").strip().upper(), ""
    return mapped


def amount_for_plan(plan: str) -> Decimal:
    return PLAN_AMOUNTS.get((plan or "").strip().lower(), Decimal("0.00"))


def full_name(names: str, lastname: str) -> str:
    return " ".join(part for part in (names.strip(), lastname.strip()) if part)


def map_cobros_row(row: MemberCobrosRow) -> tuple:
    institution, bank_code, institution_type = resolve_bank(row.member_bank_name)
    account_type, account_code = resolve_account_type(row.member_account_type)
    return (
        "",
        IDENTIFICATION_TYPE,
        (row.identifier or "").strip(),
        full_name(row.names, row.lastname),
        (row.email or "").strip(),
        institution_type,
        institution,
        account_type,
        (row.member_account_number or "").strip(),
        COBROS_REFERENCE,
        amount_for_plan(row.member_debit_plan),
        bank_code,
        account_code,
    )


def _apply_column_widths(sheet: Worksheet) -> None:
    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_cobros_header(sheet: Worksheet) -> None:
    sheet.merge_cells("B1:J1")
    sheet.merge_cells("B2:J2")
    sheet.merge_cells("B3:J3")
    sheet.merge_cells("B6:J6")
    sheet["B1"] = COBROS_TITLE
    sheet["B1"].font = TITLE_FONT
    sheet["B1"].alignment = CENTER
    sheet["B2"] = COBROS_SUBTITLE
    sheet["B2"].font = TITLE_FONT
    sheet["B2"].alignment = CENTER
    sheet["D4"] = "TOTAL REGISTROS:"
    sheet["D4"].font = LABEL_FONT
    sheet["D4"].alignment = LEFT
    sheet["E4"].font = LABEL_FONT
    sheet["E4"].number_format = "#,##0"
    sheet["G4"] = "MONTO TOTAL:"
    sheet["G4"].font = LABEL_FONT
    sheet["G4"].alignment = LEFT
    sheet["H4"].font = LABEL_FONT
    sheet["H4"].number_format = '"$"#,##0.00'
    for index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=7, column=index, value=header)
        cell.font = LABEL_FONT
        cell.alignment = CENTER


def _write_data_row(sheet: Worksheet, excel_row: int, values: tuple) -> None:
    for index, value in enumerate(values, start=1):
        cell = sheet.cell(row=excel_row, column=index, value=value if value not in (None, "") else None)
        cell.font = DATA_FONT
        if index in {3, 9, 13}:
            cell.number_format = "@"
            if value not in (None, ""):
                cell.value = str(value)
        elif index == 11:
            cell.number_format = "#,##0.00"
            cell.value = float(value) if value not in (None, "") else 0
        elif index == 12 and value is not None:
            cell.value = int(value)


def _write_bd_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("bd")
    sheet.sheet_state = "hidden"
    headers = (
        "Tipo de documentos",
        "Tipo de institución",
        "Banco",
        "Cooperativa",
        "Otras",
        "INSTITUCIÓN",
        "CÓDIGO",
        "TIPO",
        "CUENTA",
        "TIPO",
    )
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=index, value=header).font = LABEL_FONT
    sheet["A2"] = "C"
    sheet["A3"] = "P"
    sheet["A4"] = "R"
    sheet["B2"] = "Banco"
    sheet["B3"] = "Cooperativa"
    sheet["B4"] = "Otras"
    sheet["I2"] = "AHORROS"
    sheet["J2"] = "02"
    sheet["I3"] = "CORRIENTE"
    sheet["J3"] = "01"
    unique: list[tuple[str, int, str]] = []
    seen: set[int] = set()
    for institution, code, kind in CASH_BANKS.values():
        if code in seen:
            continue
        seen.add(code)
        unique.append((institution, code, kind))
    unique.sort(key=lambda item: item[0])
    tipo_code = {"Banco": 1, "Cooperativa": 2, "Otras": 3}
    for offset, (institution, code, kind) in enumerate(unique):
        excel_row = offset + 2
        sheet.cell(row=excel_row, column=6, value=institution)
        sheet.cell(row=excel_row, column=7, value=code)
        sheet.cell(row=excel_row, column=8, value=tipo_code.get(kind, ""))


def build_cobros_workbook(rows: list[MemberCobrosRow]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cobros"
    _apply_column_widths(sheet)
    _write_cobros_header(sheet)

    mapped = [map_cobros_row(row) for row in rows]
    total_amount = Decimal("0.00")
    for index, values in enumerate(mapped):
        _write_data_row(sheet, 8 + index, values)
        total_amount += Decimal(str(values[10]))

    sheet["E4"] = len(mapped)
    sheet["H4"] = float(total_amount)
    _write_bd_sheet(workbook)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
