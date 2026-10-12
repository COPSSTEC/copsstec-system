from datetime import date
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.modules.auth.domain.entities import User
from app.modules.auth.presentation.api.dependencies import get_current_user
from app.modules.members.application.use_cases import ExportMembersCobrosUseCase
from app.modules.members.domain.entities import MemberCobrosRow
from app.modules.members.infrastructure.cobros_excel import (
    COBROS_REFERENCE,
    COBROS_SUBTITLE,
    COBROS_TITLE,
    HEADERS,
    build_cobros_workbook,
    map_cobros_row,
)
from app.modules.members.presentation.api.dependencies import get_export_members_cobros_use_case


class FakeCobrosRepository:
    def __init__(self, rows: list[MemberCobrosRow] | None = None) -> None:
        self.rows = list(rows or [])

    def list_cobros_export_rows(self) -> list[MemberCobrosRow]:
        return list(self.rows)


def _row(**overrides: object) -> MemberCobrosRow:
    data = dict(
        identifier="1710034065",
        names="Ana",
        lastname="Pérez",
        email="ana@test.com",
        member_bank_name="Banco Pichincha",
        member_account_type="Ahorros",
        member_account_number="2200123456",
        member_debit_plan="monthly",
    )
    data.update(overrides)
    return MemberCobrosRow(**data)  # type: ignore[arg-type]


def _user(*, roles: list[str], user_id: int = 1) -> User:
    return User(
        id=user_id,
        name="Usuario",
        email="user@test.com",
        password_hash="x",
        state_id=1,
        email_verified_at=None,
        last_conexion=None,
        roles=roles,
        profile=None,
    )


def test_cobros_excel_maps_pichincha_ahorros_monthly() -> None:
    values = map_cobros_row(_row())
    assert values[0] == ""
    assert values[1] == "C"
    assert values[2] == "1710034065"
    assert values[3] == "Ana Pérez"
    assert values[4] == "ana@test.com"
    assert values[5] == "Banco"
    assert values[6] == "BANCO PICHINCHA"
    assert values[7] == "AHORROS"
    assert values[8] == "2200123456"
    assert values[9] == COBROS_REFERENCE
    assert values[10] == 10
    assert values[11] == 10
    assert values[12] == "02"


def test_cobros_excel_skips_unknown_bank_and_empty_plan() -> None:
    values = map_cobros_row(
        _row(
            member_bank_name="Banco Inventado",
            member_account_type="Ahorros",
            member_debit_plan="",
        ),
    )
    assert values[5] == ""
    assert values[6] == "Banco Inventado"
    assert values[11] is None
    assert values[10] == 0


def test_export_use_case_totals_and_empty_workbook() -> None:
    empty_content, empty_name = ExportMembersCobrosUseCase(FakeCobrosRepository()).execute(
        today=date(2026, 10, 12),
    )
    assert empty_name == "Plantilla_Cobros_COPSSTEC_2026-10-12.xlsx"
    empty_book = load_workbook(BytesIO(empty_content))
    cobros = empty_book["Cobros"]
    assert cobros["B1"].value == COBROS_TITLE
    assert cobros["B2"].value == COBROS_SUBTITLE
    assert [cobros.cell(row=7, column=index).value for index in range(1, 14)] == list(HEADERS)
    assert cobros["E4"].value == 0
    assert cobros["H4"].value == 0
    assert cobros["A8"].value is None
    assert empty_book["bd"].sheet_state == "hidden"

    rows = [
        _row(),
        _row(
            identifier="0102030405",
            names="Luis",
            lastname="Mora",
            email="luis@test.com",
            member_bank_name="Cooperativa JEP",
            member_account_type="Corriente",
            member_account_number="001122",
            member_debit_plan="quarterly",
        ),
    ]
    content, _filename = ExportMembersCobrosUseCase(FakeCobrosRepository(rows)).execute(
        today=date(2026, 10, 12),
    )
    book = load_workbook(BytesIO(content))
    sheet = book["Cobros"]
    assert sheet["E4"].value == 2
    assert sheet["H4"].value == 40
    assert sheet["A8"].value is None
    assert sheet["J8"].value == COBROS_REFERENCE
    assert sheet["C8"].value == "1710034065"
    assert sheet["K8"].value == 10
    assert sheet["L8"].value == 10
    assert sheet["M8"].value == "02"
    assert sheet["G9"].value == "COAC JUVENTUD ECUATORIANA PROGRESISTA LTDA."
    assert sheet["F9"].value == "Cooperativa"
    assert sheet["L9"].value == 213
    assert sheet["H9"].value == "CORRIENTE"
    assert sheet["M9"].value == "01"
    assert sheet["K9"].value == 30


def test_export_cobros_forbidden_for_member() -> None:
    app.dependency_overrides[get_current_user] = lambda: _user(roles=["miembro"])
    app.dependency_overrides[get_export_members_cobros_use_case] = lambda: ExportMembersCobrosUseCase(
        FakeCobrosRepository(),
    )
    try:
        client = TestClient(app)
        response = client.get("/api/members/export/cobros")
        assert response.status_code == 403
    finally:
        app.dependency_overrides.clear()


def test_export_cobros_returns_xlsx_for_admin() -> None:
    app.dependency_overrides[get_current_user] = lambda: _user(roles=["admin"])
    app.dependency_overrides[get_export_members_cobros_use_case] = lambda: ExportMembersCobrosUseCase(
        FakeCobrosRepository([_row()]),
    )
    try:
        client = TestClient(app)
        response = client.get("/api/members/export/cobros")
        assert response.status_code == 200
        assert "spreadsheetml" in response.headers["content-type"]
        assert "Plantilla_Cobros_COPSSTEC_" in response.headers["content-disposition"]
        book = load_workbook(BytesIO(response.content))
        assert book["Cobros"]["E4"].value == 1
        assert book["Cobros"]["K8"].value == 10
    finally:
        app.dependency_overrides.clear()
